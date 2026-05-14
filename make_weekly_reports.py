from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt


OUT_DIR = Path(r"D:\毕业设计\毕业设计\内容\我的一整套\周记")


def set_run_font(run, size=12, bold=False):
    run.bold = bold
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(size)


def add_paragraph(doc, text="", bold=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, bold=bold)
    return p


def build_doc(filename, week_title, main_items, problem_items, plan_items):
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(12)

    add_paragraph(doc, f"计算机科学与技术学院毕业设计（论文）周记{week_title}")
    add_paragraph(doc, "专业")
    add_paragraph(doc, "软件工程")
    add_paragraph(doc, "题目")
    add_paragraph(doc, "基于图学习的细胞疾病风险量化方法研究")
    add_paragraph(doc, "学生姓名")
    add_paragraph(doc, "冀雨杉")
    add_paragraph(doc, "学号")
    add_paragraph(doc, "22009200915")
    add_paragraph(doc, "指导")
    add_paragraph(doc, "教师")
    add_paragraph(doc, "王炳波")

    add_paragraph(doc, "1、主要工作内容和进展", bold=True)
    for item in main_items:
        add_paragraph(doc, item)

    add_paragraph(doc, "2、存在的主要问题和解决办法与思路", bold=True)
    for item in problem_items:
        add_paragraph(doc, item)

    add_paragraph(doc, "3、下周工作计划", bold=True)
    for item in plan_items:
        add_paragraph(doc, item)

    add_paragraph(doc, "4、导师意见", bold=True)
    add_paragraph(doc, "                                                       指导教师（签名）：")
    add_paragraph(doc, "说明：1. 本表每周由学生填写，导师签署意见。")
    add_paragraph(doc, "      2. 此表格作为答辩资格审查材料之一。")
    add_paragraph(doc, "      3. 文中字体采用宋体小四号，行间距为固定值20磅。")

    doc.save(str(OUT_DIR / filename))


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)

    build_doc(
        "_tmp_基于图学习的细胞疾病风险量化方法研究 第十三周周记.docx",
        "第 十三 周",
        [
            "本周结合项目文档对整体阶段进行了重新整理，明确模块一的 Seurat、scSHEFT 和 LIGER 三种跨模态映射方法已经完成，当前工作重心已经转到模块二的多组学改造。按照当前进度总览，模块二的输入路线已经确定为方法 B（scSHEFT），top-k 统一为 30，后续不再回头调整前置映射框架。",
            "围绕模块二的第一版方案，本周重点梳理了两条思路的差异：一种是先拼接 ATAC 与 RNA latent 再做 PCA 建图，另一种是把 ATAC 视角和 RNA reference 视角分别建图，再按每个细胞的 RNA 参考可靠性融合成 WNN-lite 图。结合前期结果，确认当前项目更需要的是让 RNA 信息进入上游的细胞表示和邻居结构，而不是只在 TRS 之后做解释。",
            "同时，本周对模块二的输出接口进行了整理，明确第一阶段需要先产出 RNA/ATAC 权重、融合邻居和可直接接入 SCAVENGE 的稀疏图对象，再在此基础上比较原始 TRS 与融合后 TRS 的差异。这样可以把多组学改造控制在可复现、可对比的范围内，避免一次性改动过大。",
        ],
        [
            "目前最大的难点是 RNA reference 的可靠性如何转成图权重。若权重过于依赖标签或真实一致率，容易把评估信息泄漏到方法中；若权重过弱，又无法真正体现多组学信息的增益。因此需要先保留一个轻量、可解释的权重规则，再根据结果逐步调整。",
            "另一个问题是，模块二不能和模块三混在一起写。模块二的重点是改造 TRS 的输入图和传播路径，模块三才是后面的差异表达、GO 富集和 marker 验证。当前阶段必须先把多组学融合主线跑通，再考虑下游解释。",
        ],
        [
            "下周计划继续推进 WNN-lite 的第一版实现，先完成 RNA/ATAC 权重定义、融合邻居生成和 mutualKNN 转换。",
            "同时准备好与原始 SCAVENGE TRS 的对比输出，重点查看 MONO、PLT、LYM 等目标性状是否比拼接-PCA 方案更稳定。",
        ],
    )

    build_doc(
        "_tmp_基于图学习的细胞疾病风险量化方法研究 第十四周周记.docx",
        "第 十四 周",
        [
            "本周继续按照模块二规划推进多组学融合工作，重点围绕 WNN-lite 图融合的第一版流程做整理和落地。结合 scSHEFT 已输出的共享 latent、RNA reference latent 以及每个细胞的参考质量指标，开始把 RNA 信息从“辅助说明”转成“邻居图构建”的直接输入，保证后续 SCAVENGE 传播路径能够同时反映 ATAC 信号和 RNA 参考可靠性。",
            "在具体推进上，本周进一步明确了第一版融合图的生成方式：先分别构建 ATAC 视角与 RNA 视角的 KNN，再根据 mean_similarity 和 n_selected 计算每个细胞的 RNA 权重，最后融合为 WNN-lite mutualKNN。这样做的目的不是重写整套方法，而是用一条更稳定的图结构替换原来的单一路径输入，让 TRS 计算更能体现多组学一致性。",
            "同时，本周也把模块二的第一阶段输出格式整理清楚，包括权重表、融合邻居表、诊断指标以及后续用于跑 TRS 的图对象。这样后面无论是比较原始 TRS、拼接-PCA TRS 还是 WNN-lite TRS，都能在统一接口下完成，便于直接形成阶段性结论。",
        ],
        [
            "当前仍然存在的主要问题是，融合图的效果对权重规则和 top-k 设置比较敏感，尤其是某些亚型细胞在 Gene Activity 或 RNA reference 上的可靠性不高时，容易让局部邻居结构发生偏移。后续需要通过诊断文件检查孤立点、度分布和权重分布，避免图结构失衡。",
            "另外，虽然模块二已经有了明确的方向，但这部分结果还不能过早写成最终结论。当前更适合把它表述为第一版融合流程和初步对比结果，等到与原始 SCAVENGE TRS 的正式比较完成后，再统一总结改进效果。",
        ],
        [
            "下周计划继续跑通融合图上的 SCAVENGE 计算，并整理原始 TRS 与 WNN-lite TRS 的对比结果。",
            "同时补充一版简短的阶段总结，把模块一完成、模块二启动和模块三后置验证的关系说明清楚，为后续论文正文和答辩汇报统一口径。",
        ],
    )


if __name__ == "__main__":
    main()
